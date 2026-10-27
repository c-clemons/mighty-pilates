import openpyxl, sys, json, re
from collections import defaultdict
out={}
for path in sys.argv[1:]:
    wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
    print(path, wb.sheetnames)
    for ws in wb.worksheets:
        if ws.title in('Locations','BS','SCF'): continue
        rows=list(ws.iter_rows(values_only=True))
        hdr=None
        for i,r in enumerate(rows[:8]):
            if r and sum(1 for c in r if c is not None and re.match(r'^[A-Z][a-z]{2} \d{4}$',str(c).strip()))>=3:
                hdr=i;break
        if hdr is None:
            print('  no hdr',ws.title, [r[:4] for r in rows[:6]]); continue
        cols={j:str(c).strip() for j,c in enumerate(rows[hdr]) if c is not None and re.match(r'^[A-Z][a-z]{2} \d{4}$',str(c).strip())}
        for r in rows[hdr+1:]:
            if not r or r[0] is None: continue
            lab=str(r[0]).strip()
            if re.match(r'^(Total )?(for )?6010\d\d',lab) or lab.startswith('Total Income') or re.match(r'^(Total )?(for )?601000',lab) or lab in('Total for Income','NET INCOME','Net Income','Total Expenses','Total for Expenses'):
                for j,m in cols.items():
                    v=r[j] if j<len(r) else None
                    out.setdefault(ws.title,{}).setdefault(lab,{})[m]=float(v) if isinstance(v,(int,float)) else 0.0
json.dump(out,open(sys.argv[0].replace('mkt.py','mkt.json'),'w'),indent=1)
for t in out:
    print(t, list(out[t].keys()))
